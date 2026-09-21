import os
from pathlib import Path
import structlog
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    ListFlowable,
    ListItem,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from app.services.resume.builder import MasterResumeData

logger = structlog.get_logger(__name__)


class ResumeExporters:
    """
    Exporters for ATS-compliant PDF, Microsoft Word DOCX, and Markdown resumes.
    """

    @staticmethod
    def export_markdown(data: MasterResumeData, output_path: str, content: str | None = None) -> str:
        """Write Markdown file to disk."""
        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        if content is None:
            from app.services.resume.builder import get_master_resume_builder
            content = get_master_resume_builder().render_markdown(data)
        p.write_text(content, encoding="utf-8")
        return str(p.absolute())

    @staticmethod
    def export_docx(data: MasterResumeData, output_path: str) -> str:
        """Export resume to ATS-friendly Microsoft Word DOCX."""
        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Set 0.5 inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # 1. Header (Name)
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(2)
        run_name = title_p.add_run(data.candidate_name)
        run_name.font.size = Pt(18)
        run_name.font.bold = True
        run_name.font.color.rgb = RGBColor(15, 23, 42)  # slate-900

        # Contact Info Line
        contact_p = doc.add_paragraph()
        contact_p.paragraph_format.space_after = Pt(8)
        contact_text = (
            f"{data.location} | {data.phone} | {data.email} | "
            f"Portfolio: {data.portfolio_url} | LinkedIn: {data.linkedin_url}"
        )
        run_contact = contact_p.add_run(contact_text)
        run_contact.font.size = Pt(9)
        run_contact.font.color.rgb = RGBColor(71, 85, 105)  # slate-600

        # Helper for Section Heading
        def add_heading(title: str):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(8)
            hp.paragraph_format.space_after = Pt(2)
            run = hp.add_run(title.upper())
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)  # blue-900

        # 2. Professional Summary
        add_heading("Professional Summary")
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        r_sum = sp.add_run(data.summary)
        r_sum.font.size = Pt(9.5)

        # 3. Technical Skills
        add_heading("Technical Skills")
        sk_p = doc.add_paragraph()
        sk_p.paragraph_format.space_after = Pt(6)
        r_prio = sk_p.add_run(f"Core Priority Focus: {', '.join(data.priority_skills)}\n")
        r_prio.font.size = Pt(9.5)
        r_prio.font.bold = True

        for cat, sk_list in data.categorized_skills.items():
            r_cat = sk_p.add_run(f"{cat}: ")
            r_cat.font.size = Pt(9)
            r_cat.font.bold = True
            r_items = sk_p.add_run(f"{', '.join(sk_list)}\n")
            r_items.font.size = Pt(9)

        # 4. Work Experience
        add_heading("Professional Experience")
        for exp in data.experiences:
            exp_header = doc.add_paragraph()
            exp_header.paragraph_format.space_before = Pt(4)
            exp_header.paragraph_format.space_after = Pt(1)
            
            r_role = exp_header.add_run(f"{exp.role} ")
            r_role.font.bold = True
            r_role.font.size = Pt(10)

            r_comp = exp_header.add_run(f"— {exp.company}")
            r_comp.font.size = Pt(10)

            r_meta = exp_header.add_run(f"  ({exp.location} | {exp.period})")
            r_meta.font.size = Pt(8.5)
            r_meta.font.italic = True
            r_meta.font.color.rgb = RGBColor(100, 116, 139)

            for b in exp.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                r_b = bp.add_run(b.text)
                r_b.font.size = Pt(9)

        # 5. Key Projects
        add_heading("Key Projects")
        for prj in data.projects:
            prj_header = doc.add_paragraph()
            prj_header.paragraph_format.space_before = Pt(4)
            prj_header.paragraph_format.space_after = Pt(1)

            r_pname = prj_header.add_run(f"{prj.name} ")
            r_pname.font.bold = True
            r_pname.font.size = Pt(10)

            r_ptitle = prj_header.add_run(f"— {prj.title}")
            r_ptitle.font.size = Pt(9.5)
            r_ptitle.font.italic = True

            r_ptech = prj_header.add_run(f"  [Tech: {', '.join(prj.technologies[:6])}]")
            r_ptech.font.size = Pt(8.5)
            r_ptech.font.color.rgb = RGBColor(71, 85, 105)

            for b in prj.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                r_b = bp.add_run(b.text)
                r_b.font.size = Pt(9)

        # 6. Education
        add_heading("Education")
        for edu in data.education:
            ed_p = doc.add_paragraph()
            ed_p.paragraph_format.space_before = Pt(2)
            ed_p.paragraph_format.space_after = Pt(2)
            r_deg = ed_p.add_run(f"{edu.degree} — {edu.institution}")
            r_deg.font.bold = True
            r_deg.font.size = Pt(9.5)
            r_emeta = ed_p.add_run(f" ({edu.period} | {edu.grade})")
            r_emeta.font.size = Pt(8.5)
            r_emeta.font.italic = True

        doc.save(str(p.absolute()))
        return str(p.absolute())

    @staticmethod
    def export_pdf(data: MasterResumeData, output_path: str) -> str:
        """Export resume to clean, professional ATS-friendly PDF using ReportLab."""
        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)

        doc = SimpleDocTemplate(
            str(p.absolute()),
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        # Custom Typography Styles
        title_style = ParagraphStyle(
            "ResumeTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=18,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=3,
        )

        contact_style = ParagraphStyle(
            "ResumeContact",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#475569"),
            spaceAfter=6,
        )

        heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor("#1e3a8a"),
            spaceBefore=6,
            spaceAfter=2,
        )

        body_style = ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=3,
        )

        item_title_style = ParagraphStyle(
            "ItemTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=12,
            textColor=colors.HexColor("#0f172a"),
        )

        meta_style = ParagraphStyle(
            "ItemMeta",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#64748b"),
        )

        bullet_style = ParagraphStyle(
            "BulletStyle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.8,
            leading=11.5,
            leftIndent=12,
            firstLineIndent=-8,
            spaceAfter=2,
            textColor=colors.HexColor("#1e293b"),
        )

        story = []

        # 1. Header (Name & Contact)
        story.append(Paragraph(data.candidate_name, title_style))
        contact_html = (
            f"{data.location} &nbsp;|&nbsp; {data.phone} &nbsp;|&nbsp; "
            f"<a href='mailto:{data.email}'>{data.email}</a> &nbsp;|&nbsp; "
            f"<a href='{data.portfolio_url}'>Portfolio</a> &nbsp;|&nbsp; "
            f"<a href='{data.linkedin_url}'>LinkedIn</a> &nbsp;|&nbsp; "
            f"<a href='{data.github_url}'>GitHub</a>"
        )
        story.append(Paragraph(contact_html, contact_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=5))

        # 2. Professional Summary
        story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
        story.append(Paragraph(data.summary, body_style))

        # 3. Technical Skills
        story.append(Paragraph("TECHNICAL SKILLS", heading_style))
        skills_text = f"<b>Core Focus:</b> {', '.join(data.priority_skills)}<br/>"
        skills_text += "<br/>".join(
            f"<b>{cat}:</b> {', '.join(s)}"
            for cat, s in data.categorized_skills.items()
        )
        story.append(Paragraph(skills_text, body_style))

        # 4. Work Experience
        story.append(Paragraph("PROFESSIONAL EXPERIENCE", heading_style))
        for exp in data.experiences:
            exp_hdr = (
                f"<b>{exp.role}</b> — {exp.company} "
                f"<font color='#64748b' size='8'><i>({exp.location} | {exp.period})</i></font>"
            )
            story.append(Paragraph(exp_hdr, item_title_style))
            for b in exp.bullets:
                story.append(Paragraph(f"• {b.text}", bullet_style))
            story.append(Spacer(1, 3))

        # 5. Key Projects
        story.append(Paragraph("KEY PROJECTS", heading_style))
        for prj in data.projects:
            prj_hdr = (
                f"<b>{prj.name}</b> — <i>{prj.title}</i> "
                f"<font color='#64748b' size='8'>[Tech: {', '.join(prj.technologies[:6])}]</font>"
            )
            story.append(Paragraph(prj_hdr, item_title_style))
            for b in prj.bullets:
                story.append(Paragraph(f"• {b.text}", bullet_style))
            story.append(Spacer(1, 3))

        # 6. Education
        story.append(Paragraph("EDUCATION", heading_style))
        for edu in data.education:
            edu_txt = f"<b>{edu.degree}</b> — {edu.institution} <font color='#64748b' size='8'><i>({edu.period} | {edu.grade})</i></font>"
            story.append(Paragraph(edu_txt, body_style))

        doc.build(story)
        return str(p.absolute())
